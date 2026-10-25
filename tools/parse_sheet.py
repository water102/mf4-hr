import openpyxl, json, sys, os
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'f:\ws\projects\hrm-mf14\data\DANH SÁCH NHÂN VIÊN MS14_new.xlsx', data_only=True)
sheet = wb['Nhan vien']

employees = []
for r in range(8, 45):
    stt = sheet.cell(r, 1).value
    name = sheet.cell(r, 5).value # Col 5 is Họ và tên
    if not name and not stt:
        continue
    
    def val(c):
        v = sheet.cell(r, c).value
        if hasattr(v, 'strftime'):
            return v.strftime('%d/%m/%Y')
        if isinstance(v, float) and v.is_integer():
            return int(v)
        return v

    code_val = val(3) or f'MS{stt:03d}' # Col 3 is Đánh lại mã số NV
    clean_code = str(code_val).strip()
    dept = val(2)
    if dept == 'BGH': dept = 'Ban giám hiệu'
    if not dept: dept = 'Vận hành'
    
    title = val(4) or 'Nhân viên' # Col 4 is Chức danh
    phone = str(val(6)) if val(6) is not None else '' # Col 6 is SĐT
    dob = val(7) or '' # Col 7 is Ngày sinh
    id_card = str(val(8)) if val(8) is not None else '' # Col 8 is CCCD/hộ chiếu
    id_date = val(9) or '' # Col 9 is Ngày cấp
    id_place = val(10) or '' # Col 10 is Nơi cấp
    hktt = val(11) or '' # Col 11 is HKTT
    address = val(12) or hktt or '' # Col 12 is Địa chỉ
    
    # TV (Cols 13-17)
    hd_tv_num = val(13)
    hd_tv_start = val(14)
    hd_tv_end = val(15)
    hd_tv_quit = val(16)
    hd_tv_remain = val(17)
    
    # CT (Cols 18-22)
    hd_ct_num = val(18)
    hd_ct_start = val(19)
    hd_ct_end = val(20)
    hd_ct_quit = val(21)
    hd_ct_remain = val(22)
    
    note = val(23) or ''
    so_yeu = val(25)
    doc_cccd = val(26)
    gks = val(27)
    bang_cap = val(28)
    gksk = val(29)
    tu_phap = val(30)
    
    net = val(32)
    bhxh_base = val(33)
    bhxh_deduct = val(34)
    gross = val(35)
    probation_salary = val(36)
    bank_acc = str(val(40)) if val(40) is not None else ''
    bank_name = val(41) or ''
    
    clean_name = str(name).strip()
    safe_name = "".join(c for c in clean_name if c.isalnum())
    
    employees.append({
        'stt': stt,
        'id': clean_code,
        'name': clean_name,
        'dept': dept,
        'position': title,
        'jobTitle': title,
        'phone': phone,
        'dob': dob,
        'idCard': id_card,
        'idDate': id_date,
        'idPlace': id_place or 'Cục Cảnh sát QLHC về TTXH',
        'hometown': id_place or ('Hà Nội' if 'Hà Nội' in hktt else hktt or 'Hà Nội'),
        'address': address,
        'hktt': hktt,
        'email': f'{clean_name.split()[-1].lower()}.{clean_code.lower()}@mf14.vn',
        'gender': 'Nữ' if any(w in clean_name.lower() for w in ['thị', 'hồng', 'hoa', 'thảo', 'hường', 'liên', 'hạnh', 'nga', 'quỳnh', 'loan', 'phương', 'mai', 'hảo', 'nhung', 'dương', 'duyên', 'thủy']) else 'Nam',
        'status': 'Đã nghỉ việc' if (hd_ct_quit or hd_tv_quit) else 'Đang làm việc',
        'leaveDate': hd_ct_quit or hd_tv_quit or '-',
        'startDate': hd_ct_start or hd_tv_start or '01/01/2026',
        'contract': {
            'number': str(hd_ct_num or hd_tv_num or f'HĐ-MF14-{clean_code}/2026'),
            'type': 'Chính thức (1 năm)' if hd_ct_num or hd_ct_start else ('Thử việc' if hd_tv_start or hd_tv_num else 'Hợp đồng lao động'),
            'startDate': hd_ct_start or hd_tv_start or '01/01/2026',
            'endDate': hd_ct_end or hd_tv_end or '31/12/2026',
            'salary': f'{int(gross):,} đ'.replace(',', '.') if gross else (f'{int(probation_salary):,} đ'.replace(',', '.') if probation_salary else (f'{int(net):,} đ (Net)'.replace(',', '.') if net else 'Thoả thuận')),
            'driveFile': f'HD_{safe_name}_Scan.pdf',
            'status': 'Sắp hết hạn' if ('hết hạn' in str(note).lower() or (hd_tv_remain and isinstance(hd_tv_remain, (int, float)) and hd_tv_remain <= 30)) else ('Chưa ký HĐ' if 'chưa ký' in str(note).lower() else 'Đang hiệu lực')
        },
        'salary': {
            'gross': gross or 0,
            'probation': probation_salary or 0,
            'net': net or 0,
            'bhxhBase': bhxh_base or 0,
            'bhxhDeduct': bhxh_deduct or 0,
            'bank': bank_name,
            'bankAccount': bank_acc
        },
        'documents': {
            'soYeuLyLich': bool(so_yeu),
            'cccd': bool(doc_cccd),
            'gks': bool(gks),
            'bangCap': bool(bang_cap),
            'gksk': bool(gksk),
            'tuPhap': bool(tu_phap)
        },
        'note': str(note).strip() if note else ''
    })

print(f'Parsed {len(employees)} employees.')
with open(r'f:\ws\projects\hrm-mf14\data\employees.json', 'w', encoding='utf-8') as f:
    json.dump(employees, f, ensure_ascii=False, indent=2)
print('Saved to f:\\ws\\projects\\hrm-mf14\\data\\employees.json')
